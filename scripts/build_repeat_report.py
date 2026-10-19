# -*- coding: utf-8 -*-
"""生成《历年重复单词汇总》Excel 报告（数据源：source/真题完形补文词汇总表.docx）。"""
from __future__ import annotations

import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from build_data import DEFAULT_SRC, parse_docx  # noqa: E402

OUT = ROOT / "reports" / "历年重复单词汇总.xlsx"

ACCENT = "A8432C"
INK = "211F1B"
MUTED = "857B6C"
LINE = "E3DAC8"
CREAM = "FBF6EC"
GREEN = "3C5F4D"

THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
ALT_FILL = PatternFill("solid", fgColor=CREAM)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
FONT = "微软雅黑"


def norm(value: str) -> str:
    return (value or "").strip().lower()


def unique_in_order(values):
    seen = set()
    out = []
    for v in values:
        if v not in seen:
            seen.add(v)
            out.append(v)
    return out


def main() -> None:
    data = parse_docx(DEFAULT_SRC)
    papers = data["papers"]
    paper_order = {p["id"]: i for i, p in enumerate(papers)}

    base_map: dict[str, list[tuple[dict, dict]]] = defaultdict(list)
    answer_map: dict[str, list[tuple[dict, dict]]] = defaultdict(list)
    pair_map: dict[tuple[str, str], list[tuple[dict, dict]]] = defaultdict(list)

    for p in papers:
        for it in p["items"]:
            b, a = norm(it["base"]), norm(it["answer"])
            base_map[b].append((p, it))
            answer_map[a].append((p, it))
            pair_map[(b, a)].append((p, it))

    def by_paper_count(rows):
        return len({p["id"] for p, _ in rows})

    def sorted_groups(mapping):
        groups = []
        for key, rows in mapping.items():
            ps = {p["id"] for p, _ in rows}
            if len(ps) >= 2:
                groups.append((key, rows, len(ps)))
        groups.sort(key=lambda x: (-x[2], -len(x[1]), str(x[0])))
        return groups

    base_groups = sorted_groups(base_map)
    answer_groups = sorted_groups(answer_map)
    pair_groups = sorted_groups(pair_map)

    # 词族：同一考点的原词与答案连边，取连通分量
    parent: dict[str, str] = {}

    def find(x: str) -> str:
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a: str, b: str) -> None:
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[ra] = rb

    for p in papers:
        for it in p["items"]:
            b, a = norm(it["base"]), norm(it["answer"])
            find(b)
            find(a)
            union(b, a)

    comps: dict[str, dict] = {}
    for p in papers:
        for it in p["items"]:
            root = find(norm(it["base"]))
            c = comps.setdefault(root, {"papers": set(), "items": [], "forms": set()})
            c["papers"].add(p["id"])
            c["items"].append((p, it))
            c["forms"].add(norm(it["base"]))
            c["forms"].add(norm(it["answer"]))

    family_groups = []
    for c in comps.values():
        if len(c["papers"]) >= 2:
            family_groups.append(c)
    family_groups.sort(key=lambda c: (-len(c["papers"]), -len(c["items"]), " / ".join(sorted(c["forms"]))))

    # 近似重复卷：完全同题集合的 Jaccard
    pair_sets = {
        p["id"]: {(norm(it["base"]), norm(it["answer"])) for it in p["items"]}
        for p in papers
    }
    near_pairs = []
    for i, a in enumerate(papers):
        for b in papers[i + 1:]:
            sa, sb = pair_sets[a["id"]], pair_sets[b["id"]]
            inter = sa & sb
            union_size = len(sa | sb)
            if len(inter) >= 5 and union_size and len(inter) / union_size >= 0.8:
                near_pairs.append((a, b, inter, len(inter), union_size))

    # 统计分布
    def dist(groups, max_n=6):
        c = Counter(n for _, _, n in groups)
        return {n: c.get(n, 0) for n in range(max_n, 1, -1)}

    base_dist = dist(base_groups)
    answer_dist = dist(answer_groups)
    pair_dist = dist(pair_groups)
    family_dist = Counter(len(c["papers"]) for c in family_groups)

    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GREEN
    widths = {"A": 26, "B": 10, "C": 18, "D": 18, "E": 18}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:E1")
    ws["A1"] = "历年重复单词汇总"
    ws["A1"].font = Font(name=FONT, size=18, bold=True, color=ACCENT)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:E2")
    ws["A2"] = "数据范围：2013年10月—2026年4月 · 26 套真题 · 261 个考点 · 按不同考期计数"
    ws["A2"].font = Font(name=FONT, size=9, color=MUTED)
    ws.row_dimensions[2].height = 18

    ws["A4"] = "关键统计"
    ws["A4"].font = Font(name=FONT, size=12, bold=True, color=ACCENT)
    ws.row_dimensions[4].height = 22

    stats_headers = ["指标", "数量", "说明"]
    for j, h in enumerate(stats_headers, 1):
        cell = ws.cell(row=6, column=j, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.merge_cells("C6:E6")

    stats = [
        ("重复原词（≥2 个考期）", len(base_groups), "同一原词在不同考期出现"),
        ("其中 ≥3 个考期", sum(1 for _, _, n in base_groups if n >= 3), "重点复习对象"),
        ("重复答案词（≥2 个考期）", len(answer_groups), "同一答案形式在不同考期出现"),
        ("完全相同考点（原词+答案）", len(pair_groups), "忽略句首大小写"),
        ("重复词族（≥2 个考期）", len(family_groups), "同一词根的不同词形合并"),
    ]
    for i, (label, value, note) in enumerate(stats, start=7):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)
        ws.cell(row=i, column=3, value=note)
        ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=5)
        for col in range(1, 6):
            cell = ws.cell(row=i, column=col)
            cell.border = BORDER
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.alignment = LEFT if col != 2 else CENTER
        ws.row_dimensions[i].height = 22

    ws["A13"] = "按出现考期数分布"
    ws["A13"].font = Font(name=FONT, size=12, bold=True, color=ACCENT)
    ws.row_dimensions[13].height = 22

    dist_headers = ["出现考期数", "原词", "答案词", "完全同题", "词族"]
    for j, h in enumerate(dist_headers, 1):
        cell = ws.cell(row=15, column=j, value=h)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER

    for i, n in enumerate(range(5, 1, -1), start=16):
        values = [n, base_dist.get(n, 0), answer_dist.get(n, 0), pair_dist.get(n, 0), family_dist.get(n, 0)]
        for j, value in enumerate(values, 1):
            cell = ws.cell(row=i, column=j, value=value)
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.alignment = CENTER
            cell.border = BORDER
            if i % 2 == 0:
                cell.fill = ALT_FILL

    notes = [
        ("原词重复", "同一个原词出现在 ≥2 个不同考期，按考期数计数；同一考期内不重复计。"),
        ("答案词重复", "同一个答案形式出现在 ≥2 个不同考期；difference 来自 different ×2 与 differ ×1。"),
        ("完全同题", "原词与答案都相同；忽略句首大小写，如 2018年10月 Probably 与 2019/2020年 probably 视为同题。"),
        ("词族", "把每个考点的原词与答案连边，取连通分量；例如 nature / natural / naturally 属于同一词族。"),
        ("2018年10月", "已按 Keeping Friendship Alive 41–50 更新，不再与 2018年4月重复。"),
        ("2019年10月 / 2020年4月", "两套考点 10/11 相同，属高度近似重复卷；blame、correct、hear、wise、possible、probable、success 等的 3 次计数包含这两套。belonging、try 的 2 次也仅来自这两套。"),
    ]
    row = 22
    ws.cell(row=row, column=1, value="统计口径与注意").font = Font(name=FONT, size=12, bold=True, color=ACCENT)
    row += 1
    for title, text in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value=f"{title}：{text}")
        cell.font = Font(name=FONT, size=9, color=INK)
        cell.alignment = LEFT
        cell.fill = PatternFill("solid", fgColor=CREAM if row % 2 == 0 else "FFFFFF")
        cell.border = BORDER
        ws.row_dimensions[row].height = 42
        row += 1

    # ---------- 数据表 ----------
    def add_sheet(name, title, subtitle, headers, rows, widths_map, center_cols=(), tab=ACCENT, row_height=46):
        sh = wb.create_sheet(name)
        sh.sheet_view.showGridLines = False
        sh.sheet_properties.tabColor = tab
        last_col = len(headers)
        last_letter = get_column_letter(last_col)

        sh.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
        sh.cell(row=1, column=1, value=title).font = Font(name=FONT, size=16, bold=True, color=ACCENT)
        sh.cell(row=1, column=1).alignment = Alignment(vertical="center")
        sh.row_dimensions[1].height = 30

        sh.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        sh.cell(row=2, column=1, value=subtitle).font = Font(name=FONT, size=9, color=MUTED)
        sh.row_dimensions[2].height = 18

        for j, header in enumerate(headers, 1):
            cell = sh.cell(row=4, column=j, value=header)
            cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
            cell.alignment = CENTER
            cell.border = BORDER
        sh.row_dimensions[4].height = 30

        for i, row_values in enumerate(rows, start=5):
            for j, value in enumerate(row_values, 1):
                cell = sh.cell(row=i, column=j, value=value)
                cell.font = Font(name=FONT, size=10, color=INK)
                cell.alignment = CENTER if j in center_cols else LEFT
                cell.border = BORDER
                if i % 2 == 0:
                    cell.fill = ALT_FILL
            sh.row_dimensions[i].height = row_height

        for col, width in widths_map.items():
            sh.column_dimensions[col].width = width
        sh.freeze_panes = "A5"
        if rows:
            sh.auto_filter.ref = f"A4:{last_letter}{4 + len(rows)}"
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.print_title_rows = "1:4"
        return sh

    # 高频原词
    high_rows = []
    for key, rows, count in base_groups:
        if count < 3:
            continue
        rows_sorted = sorted(rows, key=lambda x: paper_order[x[0]["id"]])
        labels = [p["label"] for p, _ in rows_sorted]
        answers = [it["answer"] for _, it in rows_sorted]
        detail = "；".join(f"{p['label']}：{it['answer']}" for p, it in rows_sorted)
        high_rows.append((rows_sorted[0][1]["base"], count, "、".join(labels), "、".join(unique_in_order(answers)), detail))

    add_sheet(
        "高频原词",
        "高频重复原词（≥3 个考期）",
        "按出现考期数排序；同一原词在不同考期的答案形式可能不同",
        ["原词", "出现考期数", "考期", "答案形式", "考期与答案"],
        high_rows,
        {"A": 16, "B": 11, "C": 40, "D": 30, "E": 85},
        center_cols={2},
    )

    # 全部原词
    all_base_rows = []
    for key, rows, count in base_groups:
        rows_sorted = sorted(rows, key=lambda x: paper_order[x[0]["id"]])
        labels = [p["label"] for p, _ in rows_sorted]
        answers = [it["answer"] for _, it in rows_sorted]
        detail = "；".join(f"{p['label']}：{it['answer']}" for p, it in rows_sorted)
        all_base_rows.append((rows_sorted[0][1]["base"], count, "、".join(labels), "、".join(unique_in_order(answers)), detail))

    add_sheet(
        "全部原词",
        "全部重复原词（≥2 个考期）",
        "共 47 个；按出现考期数排序",
        ["原词", "出现考期数", "考期", "答案形式", "考期与答案"],
        all_base_rows,
        {"A": 16, "B": 11, "C": 46, "D": 34, "E": 95},
        center_cols={2},
    )

    # 完全同题
    exact_rows = []
    for (b, a), rows, count in pair_groups:
        rows_sorted = sorted(rows, key=lambda x: paper_order[x[0]["id"]])
        labels = [p["label"] for p, _ in rows_sorted]
        base = rows_sorted[0][1]["base"]
        answer = rows_sorted[0][1]["answer"]
        note = ""
        if a == "probably" and any(it["answer"] == "Probably" for _, it in rows_sorted):
            note = "2018年10月为句首大写 Probably"
        exact_rows.append((base, answer, count, "、".join(labels), note))

    add_sheet(
        "完全同题",
        "完全相同的考点（原词 + 答案）",
        "共 37 组；忽略句首大小写；按重复次数排序",
        ["原词", "答案", "重复次数", "考期", "备注"],
        exact_rows,
        {"A": 18, "B": 22, "C": 11, "D": 62, "E": 36},
        center_cols={3},
    )

    # 答案词
    answer_rows = []
    for key, rows, count in answer_groups:
        rows_sorted = sorted(rows, key=lambda x: paper_order[x[0]["id"]])
        labels = [p["label"] for p, _ in rows_sorted]
        sources = unique_in_order([it["base"] for _, it in rows_sorted])
        detail = "；".join(f"{p['label']}：{it['base']}→{it['answer']}" for p, it in rows_sorted)
        answer_rows.append((rows_sorted[0][1]["answer"], count, "、".join(labels), "、".join(sources), detail))

    add_sheet(
        "答案词",
        "重复答案词（≥2 个考期）",
        "共 37 个；difference 来自 different ×2 与 differ ×1",
        ["答案词", "出现考期数", "考期", "来源原词", "考期与答案"],
        answer_rows,
        {"A": 18, "B": 11, "C": 50, "D": 30, "E": 90},
        center_cols={2},
    )

    # 词族
    family_rows = []
    for c in family_groups:
        count = len(c["papers"])
        ps = [p for p in papers if p["id"] in c["papers"]]
        labels = [p["label"] for p in ps]
        forms = " / ".join(sorted(c["forms"], key=lambda s: (len(s), s)))
        detail = "；".join(
            f"{p['label']}：{it['base']}→{it['answer']}"
            for p, it in sorted(c["items"], key=lambda x: paper_order[x[0]["id"]])
        )
        family_rows.append((forms, count, "、".join(labels), detail))

    add_sheet(
        "词族",
        "重复词族（同一词根的不同词形）",
        "共 47 个词族；把每个考点的原词与答案连边后合并",
        ["词族形式", "出现考期数", "考期", "涉及考点（原词→答案）"],
        family_rows,
        {"A": 36, "B": 11, "C": 60, "D": 100},
        center_cols={2},
        row_height=60,
    )

    # 近似重复卷
    near_rows = []
    for a, b, inter, common_count, union_size in near_pairs:
        common = "；".join(f"{base}→{answer}" for base, answer in sorted(inter))
        near_rows.append((a["label"], b["label"], common_count, len(a["items"]), len(b["items"]), f"{common_count}/{union_size}", common))

    add_sheet(
        "近似重复卷",
        "高度近似的重复试卷",
        "完全同题交集 ≥5 且 Jaccard ≥ 0.8；用于解释重复次数",
        ["考期 A", "考期 B", "相同考点数", "A 考点数", "B 考点数", "重合度", "相同考点（原词→答案）"],
        near_rows,
        {"A": 14, "B": 14, "C": 12, "D": 10, "E": 10, "F": 10, "G": 100},
        center_cols={3, 4, 5, 6},
        tab=GREEN,
        row_height=60,
    )

    wb.active = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"OK  {OUT}")
    print(f"base={len(base_groups)} high={len(high_rows)} exact={len(pair_groups)} answer={len(answer_groups)} family={len(family_groups)} near={len(near_rows)}")
    print("dist base", base_dist, "answer", answer_dist, "pair", pair_dist, "family", dict(family_dist))


if __name__ == "__main__":
    main()
