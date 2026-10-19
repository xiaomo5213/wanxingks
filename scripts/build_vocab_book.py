# -*- coding: utf-8 -*-
"""生成按考期分类的《真题完形补文词汇总表》DOCX（用于导出 PDF）。"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from build_data import DEFAULT_SRC, parse_docx  # noqa: E402

OUT = ROOT / "reports" / "真题完形补文词汇总表-按考期.docx"

ACCENT = "A8432C"
INK = "211F1B"
MUTED = "857B6C"
LINE = "E3DAC8"
CREAM = "FBF6EC"
GREEN = "3C5F4D"
WHITE = "FFFFFF"

SERIF_CN = "SimSun"
SANS_CN = "Microsoft YaHei"
SERIF_EN = "Georgia"
MONO = "Consolas"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, size=10, color=INK, bold=False, ascii_font=SANS_CN, east_font=SANS_CN):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_font)


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def cell_border(cell, edge, size=4, color=LINE, val="single"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)
    borders.append(el)


def table_borders(table, color=LINE, size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in ("left", "right", "insideV"):
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def set_table_width(table, width_mm):
    tbl_pr = table._tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(Mm(width_mm).twips)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def para_border(paragraph, edge="bottom", color=LINE, size=4, space=4):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    p_bdr.append(el)


def add_page_field(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), SANS_CN)
    rfonts.set(qn("w:hAnsi"), SANS_CN)
    rfonts.set(qn("w:eastAsia"), SANS_CN)
    rpr.append(rfonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "15")  # 7.5pt
    rpr.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MUTED)
    rpr.append(color)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def add_header_footer(doc, section):
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.style = doc.styles["Normal"]
    hp.paragraph_format.tab_stops.clear_all()
    hp.paragraph_format.tab_stops.add_tab_stop(Mm(174), WD_TAB_ALIGNMENT.RIGHT)
    hp.paragraph_format.space_after = Pt(2)
    run = hp.add_run("真题完形补文词汇总表")
    set_run_font(run, size=7.5, color=MUTED)
    run = hp.add_run("\t按考期分类 · 2013—2026")
    set_run_font(run, size=7.5, color=MUTED)
    para_border(hp, "bottom", LINE, 4)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.style = doc.styles["Normal"]
    fp.paragraph_format.tab_stops.clear_all()
    fp.paragraph_format.tab_stops.add_tab_stop(Mm(174), WD_TAB_ALIGNMENT.RIGHT)
    fp.paragraph_format.space_before = Pt(2)
    run = fp.add_run("自考英语（专升本） · 历年真题")
    set_run_font(run, size=7.5, color=MUTED)
    run = fp.add_run("\t第 ")
    set_run_font(run, size=7.5, color=MUTED)
    add_page_field(fp)
    run = fp.add_run(" 页")
    set_run_font(run, size=7.5, color=MUTED)
    para_border(fp, "top", LINE, 4)


def set_cell_paragraph(p, text, size, color, bold=False, ascii_font=SANS_CN, east_font=SANS_CN, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, ascii_font=ascii_font, east_font=east_font)
    return run


def main() -> None:
    data = parse_docx(DEFAULT_SRC)
    papers = data["papers"]
    total_items = sum(len(p["items"]) for p in papers)
    years = sorted({p["year"] for p in papers})

    doc = Document()
    doc.core_properties.title = "真题完形补文词汇总表（按考期分类）"
    doc.core_properties.subject = "自考英语（专升本）历年真题完形补文词汇"
    doc.core_properties.author = "完形文库"
    doc.core_properties.keywords = "自考英语, 完形补文, 词汇, 真题"

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(16)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)
    add_header_footer(doc, section)

    # ---------------- 封面 ----------------
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(46)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("自考英语（专升本） · 历年真题")
    set_run_font(run, size=9.5, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    para_border(p, "bottom", ACCENT, 10)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.04
    run = p.add_run("真题完形补文")
    set_run_font(run, size=38, color=INK, bold=True, ascii_font=SERIF_EN, east_font=SERIF_CN)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.04
    run = p.add_run("词汇总表")
    set_run_font(run, size=38, color=ACCENT, bold=True, ascii_font=SERIF_EN, east_font=SERIF_CN)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("按考期分类 · 2013—2026")
    set_run_font(run, size=14, color=INK, ascii_font=SANS_CN, east_font=SANS_CN)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    run = p.add_run(f"{len(papers)} 套真题   ·   {total_items} 个考点   ·   {len(years)} 个年份")
    set_run_font(run, size=10, color=MUTED)

    note = doc.add_table(rows=1, cols=1)
    note.alignment = WD_TABLE_ALIGNMENT.LEFT
    note.autofit = False
    set_table_width(note, 150)
    cell = note.cell(0, 0)
    cell.width = Mm(150)
    shade(cell, CREAM)
    cell_margins(cell, top=110, bottom=110, left=160, right=160)
    cell_border(cell, "top", 4, LINE)
    cell_border(cell, "bottom", 4, LINE)
    cell_border(cell, "left", 4, LINE)
    cell_border(cell, "right", 4, LINE)
    np_ = cell.paragraphs[0]
    np_.paragraph_format.space_after = Pt(2)
    run = np_.add_run("每个考期按 41–50 题整理。")
    set_run_font(run, size=9.5, color=INK)
    np2 = cell.add_paragraph()
    np2.paragraph_format.space_after = Pt(0)
    run = np2.add_run("原词与答案并列；朱红色为答案词，右侧标注词形变化类型。")
    set_run_font(run, size=9.5, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(190)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("数据来源：真题完形补文词汇总表.docx")
    set_run_font(run, size=8, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("2026 年 10 月更新 · 适合 A4 打印")
    set_run_font(run, size=8, color=MUTED)

    # ---------------- 目录 ----------------
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run("目录")
    set_run_font(run, size=24, color=ACCENT, bold=True, ascii_font=SERIF_EN, east_font=SERIF_CN)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(f"{len(papers)} 个考期 · 按考试时间顺序排列")
    set_run_font(run, size=9, color=MUTED)

    toc = doc.add_table(rows=1 + len(papers), cols=3)
    toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc.autofit = False
    set_table_width(toc, 174)
    widths = (12, 23, 139)
    for idx, w in enumerate(widths):
        toc.columns[idx].width = Mm(w)
    for row in toc.rows:
        for idx, w in enumerate(widths):
            row.cells[idx].width = Mm(w)
    table_borders(toc, LINE, 4)

    headers = ("#", "考期", "试卷")
    for idx, text in enumerate(headers):
        cell = toc.rows[0].cells[idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_margins(cell, top=30, bottom=30, left=70, right=70)
        p_cell = cell.paragraphs[0]
        set_cell_paragraph(p_cell, text, 7.5, MUTED, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
        cell_border(cell, "bottom", 10, ACCENT)
    no_split(toc.rows[0])

    for i, paper in enumerate(papers, start=1):
        row = toc.rows[i]
        row.height = Mm(7.0)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        no_split(row)
        for idx in range(3):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell, top=10, bottom=10, left=70, right=70)
            if i % 2 == 0:
                shade(cell, CREAM)

        p_cell = row.cells[0].paragraphs[0]
        set_cell_paragraph(p_cell, f"{i:02d}", 7.5, ACCENT, ascii_font=MONO, east_font=SANS_CN)

        p_cell = row.cells[1].paragraphs[0]
        set_cell_paragraph(p_cell, paper["label"], 8.2, MUTED)

        p_cell = row.cells[2].paragraphs[0]
        p_cell.paragraph_format.space_before = Pt(0)
        p_cell.paragraph_format.space_after = Pt(0)
        p_cell.paragraph_format.line_spacing = 1.0
        run = p_cell.add_run(paper["en"])
        set_run_font(run, size=8.8, color=INK, ascii_font=SERIF_EN, east_font=SERIF_CN)
        run = p_cell.add_run("  ·  ")
        set_run_font(run, size=8, color=MUTED)
        run = p_cell.add_run(paper["zh"])
        set_run_font(run, size=8.2, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("图例：原词（词性）释义  →  答案（词性）释义。朱红色为答案词。")
    set_run_font(run, size=8, color=MUTED)

    # ---------------- 各考期 ----------------
    for index, paper in enumerate(papers, start=1):
        # 顶部信息行
        p = doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(Mm(174), WD_TAB_ALIGNMENT.RIGHT)
        run = p.add_run(f"{paper['label']}  ·  {paper['range']}")
        set_run_font(run, size=8.5, color=ACCENT, bold=True)
        run = p.add_run(f"\tNO. {index:02d} / {len(papers):02d}")
        set_run_font(run, size=8.2, color=MUTED, ascii_font=MONO, east_font=SANS_CN)

        title = paper["en"]
        title_size = 24
        if len(title) > 45:
            title_size = 17
        elif len(title) > 32:
            title_size = 20

        h = doc.add_paragraph(style="Heading 1")
        h.paragraph_format.space_before = Pt(0)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.line_spacing = 1.0
        run = h.add_run(title)
        set_run_font(run, size=title_size, color=INK, bold=True, ascii_font=SERIF_EN, east_font=SERIF_CN)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{paper['zh']}  ·  完形补文词汇对照  ·  共 {len(paper['items'])} 个考点")
        set_run_font(run, size=10, color=MUTED)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        para_border(p, "bottom", ACCENT, 8)

        items = paper["items"]
        table = doc.add_table(rows=1 + len(items), cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_width(table, 174)
        col_widths = (10, 66, 66, 32)
        for idx, w in enumerate(col_widths):
            table.columns[idx].width = Mm(w)
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Mm(w)
        table_borders(table, LINE, 4)

        header = table.rows[0]
        header.height = Mm(9)
        header.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        header_row(header)
        no_split(header)
        for idx, text in enumerate(("#", "原词 · 词性 · 释义", "答案 · 词性 · 释义", "变化")):
            cell = header.cells[idx]
            shade(cell, ACCENT)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell, top=40, bottom=40, left=80, right=80)
            p_cell = cell.paragraphs[0]
            set_cell_paragraph(p_cell, text, 8, WHITE, bold=True)

        for n, item in enumerate(items, start=1):
            row = table.rows[n]
            row.height = Mm(17.5 if len(items) <= 10 else 17)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            no_split(row)
            for idx in range(4):
                cell = row.cells[idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell_margins(cell, top=55, bottom=55, left=80, right=80)
                if n % 2 == 0:
                    shade(cell, CREAM)

            p_cell = row.cells[0].paragraphs[0]
            set_cell_paragraph(p_cell, f"{n:02d}", 7.5, MUTED, ascii_font=MONO, east_font=SANS_CN, align=WD_ALIGN_PARAGRAPH.CENTER)

            # 原词
            cell = row.cells[1]
            p1 = cell.paragraphs[0]
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            p1.paragraph_format.line_spacing = 1.0
            run = p1.add_run(item["base"])
            set_run_font(run, size=13, color=INK, bold=True, ascii_font=SERIF_EN, east_font=SERIF_CN)
            run = p1.add_run("  " + item["pos"])
            set_run_font(run, size=7.3, color=MUTED)
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_before = Pt(1)
            p2.paragraph_format.space_after = Pt(0)
            p2.paragraph_format.line_spacing = 1.0
            run = p2.add_run(item["gloss"])
            set_run_font(run, size=8.5, color=INK)

            # 答案
            cell = row.cells[2]
            p1 = cell.paragraphs[0]
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            p1.paragraph_format.line_spacing = 1.0
            run = p1.add_run(item["answer"])
            set_run_font(run, size=13, color=ACCENT, bold=True, ascii_font=SERIF_EN, east_font=SERIF_CN)
            run = p1.add_run("  " + item["apos"])
            set_run_font(run, size=7.3, color=MUTED)
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_before = Pt(1)
            p2.paragraph_format.space_after = Pt(0)
            p2.paragraph_format.line_spacing = 1.0
            run = p2.add_run(item["agloss"])
            set_run_font(run, size=8.5, color=INK)

            # 变化
            cell = row.cells[3]
            p_cell = cell.paragraphs[0]
            set_cell_paragraph(p_cell, " / ".join(item["tags"]), 7.8, GREEN)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"OK  {OUT}")
    print(f"papers={len(papers)} items={total_items} years={len(years)}")


if __name__ == "__main__":
    main()
